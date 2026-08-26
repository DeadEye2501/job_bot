"""Следствия чтения резюме с диска: что достаётся из файла и что мешает."""

import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from adapters.cv_text import read_cv_text


def _tiny_pdf(pages: list[list[tuple[int, str]]]) -> bytes:
    """Настоящий PDF из строк, поставленных на названную высоту страницы.

    Собран руками, а не библиотекой: писать PDF в проекте нечем, а подсунуть
    читателю заглушку значило бы проверить заглушку. Шрифт взят встроенный,
    поэтому строки латиницей — от алфавита чтение не зависит, а кириллица во
    встроенном шрифте потребовала бы таблицы перекодировки.

    Высота у каждой строки своя, и задаётся она нарочно: порядок, в котором
    строки записаны в файл, тогда можно развести с порядком, в котором они
    стоят на странице, — а именно этим настоящие резюме и отличаются.
    """
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [{}] /Count {} >>".format(
            " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages))), len(pages)
        ).encode(),
    ]
    for number, lines in enumerate(pages):
        drawing = "BT /F1 12 Tf\n"
        for height, line in lines:
            # Скобки и косая — разметка потока содержимого. Не заэкранируй их,
            # и строка со скобкой не уронит сборку, а молча испортит поток:
            # тест тогда соврёт вместо того, чтобы упасть.
            written = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            drawing += f"1 0 0 1 50 {height} Tm ({written}) Tj\n"
        drawing += "ET"
        body = drawing.encode("ascii")
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                "/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 "
                f"/BaseFont /Helvetica >> >> >> /Contents {4 + number * 2} 0 R >>"
            ).encode()
        )
        objects.append(
            b"<< /Length "
            + str(len(body)).encode()
            + b" >>\nstream\n"
            + body
            + b"\nendstream"
        )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    table = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{table}\n%%EOF\n".encode()
    )
    return out.getvalue()


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_plain_file_gives_back_what_is_written_in_it(tmp_path, suffix):
    resume = tmp_path / f"Резюме{suffix}"
    resume.write_text("Петров Иван\nPython Backend Developer", encoding="utf-8")

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "Петров Иван" in text
    assert "Python Backend Developer" in text


def test_pdf_gives_the_text_of_every_page(tmp_path):
    resume = tmp_path / "Резюме.pdf"
    resume.write_bytes(_tiny_pdf([[(700, "Frolov Alexey")], [(700, "Opyt raboty")]]))

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "Frolov Alexey" in text
    assert "Opyt raboty" in text


def test_pdf_lines_come_in_the_order_they_stand_on_the_page(tmp_path):
    """Порядок строк на странице, а не порядок их записи в файл.

    В живом резюме эти два порядка расходятся: заголовок «Достижения» лежит в
    файле раньше своих пунктов, а вот сами пункты соседних работ идут вперемешку
    и оказываются под чужим заголовком. Прочитанное так резюме уже не починит
    никакая модель — связи работы с её пунктами в тексте попросту нет.
    """
    resume = tmp_path / "Резюме.pdf"
    resume.write_bytes(
        _tiny_pdf([[(600, "Obyazannosti"), (700, "Dostizheniya")]])
    )

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert text.index("Dostizheniya") < text.index("Obyazannosti")


def test_docx_gives_the_text_of_its_paragraphs(tmp_path):
    resume = tmp_path / "Резюме.docx"
    document = Document()
    document.add_paragraph("Петров Иван")
    document.add_paragraph("Python Backend Developer")
    document.save(resume)

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "Петров Иван" in text
    assert "Python Backend Developer" in text


def test_docx_gives_the_text_written_inside_a_table(tmp_path):
    """Резюме сплошь и рядом свёрстано таблицей, и весь опыт лежит в ячейках."""
    resume = tmp_path / "Резюме.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "2020 — 2024"
    table.cell(0, 1).text = "Ведущий разработчик"
    document.save(resume)

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "2020 — 2024" in text
    assert "Ведущий разработчик" in text


def test_docx_keeps_paragraphs_and_tables_in_the_order_of_the_document(tmp_path):
    """Абзацы и таблицы чередуются, и порядок между ними — тот, что в файле.

    Ради этого документ и обходится поэлементно. Прочитай сперва все абзацы, а
    потом все таблицы — и подпись под таблицей уедет к следующей работе.
    """
    resume = tmp_path / "Резюме.docx"
    document = Document()
    document.add_paragraph("Опыт работы")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Ведущий разработчик"
    document.add_paragraph("Образование")
    document.save(resume)

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert text.index("Опыт работы") < text.index("Ведущий разработчик")
    assert text.index("Ведущий разработчик") < text.index("Образование")


def test_docx_gives_the_text_wrapped_in_a_content_control(tmp_path):
    """Резюме по шаблону Word держит поля в «элементе управления содержимым».

    Такой абзац лежит не в теле напрямую, и обход, знающий только абзацы да
    таблицы, прошёл бы мимо — текст пропал бы молча, без всякой причины
    наружу.
    """
    resume = tmp_path / "Резюме.docx"
    document = Document()
    body = document.element.body
    control = body.makeelement(qn("w:sdt"), {})
    inside = body.makeelement(qn("w:sdtContent"), {})
    control.append(inside)
    body.append(control)
    inside.append(document.add_paragraph("Петров Иван")._p)
    document.save(resume)

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "Петров Иван" in text


def test_a_cell_stretched_across_columns_is_read_once(tmp_path):
    """Слитая ячейка накрывает две колонки сетки, но текст в ней один."""
    resume = tmp_path / "Резюме.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Опыт работы"
    document.save(resume)

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert text.count("Опыт работы") == 1


def test_a_text_file_in_the_windows_encoding_is_read(tmp_path):
    """Русский `.txt` с русской Windows приходит не в UTF-8, и это не беда."""
    resume = tmp_path / "Резюме.txt"
    resume.write_bytes("Петров Иван".encode("cp1251"))

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "Петров Иван" in text


def test_the_byte_order_mark_does_not_reach_the_text(tmp_path):
    """«Блокнот» ставит метку в начало файла, и пробелом она не считается."""
    resume = tmp_path / "Резюме.txt"
    resume.write_bytes(b"\xef\xbb\xbf" + "Петров Иван".encode("utf-8"))

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert text.startswith("Петров")


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_a_file_lying_about_its_format_is_told_apart(tmp_path, suffix):
    """Расширение верное, а внутри не то: чаще всего `.doc`, переименованный,
    чтобы пройти фильтр диалога. Человеку про такой файл говорят словами."""
    resume = tmp_path / f"Резюме{suffix}"
    resume.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    text, trouble = read_cv_text(resume)

    assert text is None
    assert trouble == "bad_format"


def test_the_suffix_is_read_whatever_its_case(tmp_path):
    resume = tmp_path / "Резюме.PDF"
    resume.write_bytes(_tiny_pdf([[(700, "Frolov Alexey")]]))

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "Frolov Alexey" in text


def test_a_file_that_is_not_there_is_told_apart(tmp_path):
    text, trouble = read_cv_text(tmp_path / "Резюме.pdf")

    assert text is None
    assert trouble == "no_file"


def test_the_old_binary_doc_is_refused_by_its_own_name(tmp_path):
    """Диалог выбора файла его предлагает, а прочесть его нечем."""
    resume = tmp_path / "Резюме.doc"
    resume.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    text, trouble = read_cv_text(resume)

    assert text is None
    assert trouble == "bad_format"


def test_a_pdf_without_a_text_layer_is_told_apart(tmp_path):
    """Скан вместо резюме: страницы есть, а букв в них нет."""
    resume = tmp_path / "Резюме.pdf"
    resume.write_bytes(_tiny_pdf([[]]))

    text, trouble = read_cv_text(resume)

    assert text is None
    assert trouble == "no_text"


def test_an_empty_text_file_is_told_apart(tmp_path):
    resume = tmp_path / "Резюме.txt"
    resume.write_text("   \n\n  ", encoding="utf-8")

    text, trouble = read_cv_text(resume)

    assert text is None
    assert trouble == "no_text"


def test_a_text_file_in_utf16_is_read_as_written(tmp_path):
    """«Блокнот» умеет сохранять и в UTF-16, и такой файл не мусор.

    Однобайтовая кодировка разбирает его без ошибки — в бессмыслицу с нулями
    между буквами. Молча принять её значило бы отдать модели мусор вместо
    резюме.
    """
    resume = tmp_path / "Резюме.txt"
    resume.write_bytes("Петров Иван\nPython".encode("utf-16"))

    text, trouble = read_cv_text(resume)

    assert trouble is None
    assert "Петров Иван" in text
